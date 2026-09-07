# -*- coding: utf-8 -*-
from __future__ import annotations

import os
import tempfile
import unittest

from openpyxl import load_workbook

from xnt_excel_export import XntExcelExportMixin


class XntExcelTemplateTests(unittest.TestCase):
    def _rows(self):
        return [
            {
                "productId": 101,
                "productName": "Amoxicillin 500 mg viên nang cứng quy cách hộp 10 vỉ x 10 viên",
                "unit": "Viên",
                "lotNo": "AMX-001",
                "expiryDate": "2028-12-31",
                "fundSource": "Chương trình A",
                "opening": 125.5,
                "inbound": 50,
                "outbound": 20.5,
                "closing": 155,
            },
            {
                "productId": 102,
                "productName": "Vaccine thử nghiệm",
                "unit": "Lọ",
                "lotNo": "VAC-002",
                "expiryDate": "2027-06-30",
                "fundSource": "Ngân sách",
                "opening": 10,
                "inbound": 5,
                "outbound": 3,
                "closing": 12,
            },
        ]

    @staticmethod
    def _normalize_print_title_rows(value):
        return str(value or "").replace("$", "")

    def test_workbook_is_a_printable_form_not_raw_table(self):
        wb = XntExcelExportMixin._build_xnt_workbook(
            self._rows(), "2026-07-01", "2026-07-31", "Tất cả"
        )
        ws = wb.active

        self.assertEqual(ws.title, "Báo cáo XNT")
        self.assertEqual(ws["A1"].value, "BÁO CÁO XUẤT - NHẬP - TỒN")
        self.assertIn("01-07-2026", ws["A2"].value)
        self.assertIn("31-07-2026", ws["A2"].value)
        self.assertIn("Đơn vị:", ws["A3"].value)
        self.assertEqual(ws["H6"].value, "SỐ LƯỢNG")
        self.assertEqual(ws["H7"].value, "Tồn đầu")
        self.assertEqual(ws["I7"].value, "Nhập trong kỳ")
        self.assertEqual(ws["J7"].value, "Xuất trong kỳ")
        self.assertEqual(ws["K7"].value, "Tồn cuối")

        merged = {str(rng) for rng in ws.merged_cells.ranges}
        self.assertIn("A1:K1", merged)
        self.assertIn("H6:K6", merged)
        self.assertIn("A6:A7", merged)
        self.assertIn("G6:G7", merged)

        # First data row starts below the two-level header.
        self.assertEqual(ws["A8"].value, 1)
        self.assertEqual(ws["C8"].value, self._rows()[0]["productName"])
        self.assertEqual(ws["F8"].value, "31-12-2028")
        self.assertEqual(ws["H8"].value, 125.5)
        self.assertEqual(ws["K8"].value, 155)

        total_row = 10
        self.assertEqual(ws[f"A{total_row}"].value, "TỔNG CỘNG")
        self.assertEqual(ws[f"H{total_row}"].value, "=SUM(H8:H9)")
        self.assertEqual(ws[f"K{total_row}"].value, "=SUM(K8:K9)")

        # Print contract: A4 landscape, fit to one page wide, repeat headers.
        self.assertEqual(ws.freeze_panes, "A8")
        self.assertEqual(ws.page_setup.orientation, "landscape")
        self.assertEqual(str(ws.page_setup.paperSize), str(ws.PAPERSIZE_A4))
        self.assertEqual(ws.page_setup.fitToWidth, 1)
        self.assertEqual(self._normalize_print_title_rows(ws.print_title_rows), "6:7")
        self.assertFalse(ws.sheet_view.showGridLines)
        print_area_text = " ".join(str(part) for part in ws.print_area)
        self.assertIn("$A$1", print_area_text)
        self.assertIn("$K$", print_area_text)
        self.assertIn("Trang &P / &N", ws.oddFooter.center.text)

        # The report is an operational form, not a claimed legal form number.
        all_text = " ".join(
            str(cell.value or "")
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=11)
            for cell in row
        )
        self.assertNotIn("S12-H", all_text)
        self.assertIn("NGƯỜI LẬP BIỂU", all_text)
        self.assertIn("THỦ KHO", all_text)

    def test_saved_workbook_reopens_with_layout_contract(self):
        wb = XntExcelExportMixin._build_xnt_workbook(
            self._rows(), "2026-07-01", "2026-07-31", "Nguồn A"
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            path = os.path.join(temp_dir, "xnt.xlsx")
            wb.save(path)
            reopened = load_workbook(path, data_only=False)
            ws = reopened["Báo cáo XNT"]
            self.assertEqual(ws["A1"].value, "BÁO CÁO XUẤT - NHẬP - TỒN")
            self.assertIn("Nguồn A", ws["A4"].value)
            self.assertEqual(ws["H10"].value, "=SUM(H8:H9)")
            self.assertEqual(self._normalize_print_title_rows(ws.print_title_rows), "6:7")


if __name__ == "__main__":
    unittest.main()
