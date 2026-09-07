# -*- coding: utf-8 -*-
import os
import tempfile
import unittest
from unittest import mock

from openpyxl import load_workbook

from xnt_excel_export import XntExcelExportMixin


class _DummyDB:
    def xnt_report(self, start_s, end_s, fund_source):
        return [
            {
                "productId": 101,
                "productName": "Paracetamol 500 mg viên nén - tên sản phẩm dài để kiểm tra wrap",
                "lotNo": "LOT-2026-01",
                "expiryDate": "2028-12-31",
                "fundSource": "Ngân sách nhà nước",
                "opening": 100,
                "inbound": 25,
                "outbound": 40,
                "closing": 85,
            },
            {
                "productId": 102,
                "productName": "Vaccine mẫu",
                "lotNo": "VAC-02",
                "expiryDate": "2027-06-30",
                "fundSource": "Chương trình mục tiêu",
                "opening": 10,
                "inbound": 5,
                "outbound": 3,
                "closing": 12,
            },
        ]

    def q(self, sql):
        return [
            {"id": 101, "defaultUnit": "Viên"},
            {"id": 102, "defaultUnit": "Lọ"},
        ]


class _Entry:
    def __init__(self, value):
        self.entry = self
        self._value = value

    def get(self):
        return self._value


class _Combo:
    def get(self):
        return "Tất cả"


class _App(XntExcelExportMixin):
    def __init__(self):
        self.db = _DummyDB()
        self.de_from = _Entry("01-07-2026")
        self.de_to = _Entry("31-07-2026")
        self.cmb_report_fund = _Combo()
        self.toasts = []

    def _date_range_from_entries(self, start, end):
        return "2026-07-01", "2026-07-31"

    def toast(self, message, *args, **kwargs):
        self.toasts.append(message)


class TestXntExcelForm(unittest.TestCase):
    def test_printable_form_layout_and_totals(self):
        app = _App()
        with tempfile.TemporaryDirectory() as td:
            path = os.path.join(td, "xnt.xlsx")
            with mock.patch("xnt_excel_export.filedialog.asksaveasfilename", return_value=path), \
                 mock.patch("xnt_excel_export.os.startfile", create=True), \
                 mock.patch("xnt_excel_export.messagebox.showerror") as showerror:
                app.export_report_excel()

            showerror.assert_not_called()
            self.assertTrue(os.path.exists(path))

            ws = load_workbook(path).active
            self.assertEqual(ws.title, "Báo cáo XNT")
            self.assertEqual(ws["A1"].value, "BÁO CÁO XUẤT - NHẬP - TỒN")
            self.assertIn("01-07-2026", ws["A2"].value)
            self.assertIn("31-07-2026", ws["A2"].value)
            self.assertEqual(ws["H5"].value, "SỐ LƯỢNG")
            self.assertEqual(ws["H6"].value, "Tồn đầu")
            self.assertEqual(ws["I6"].value, "Nhập trong kỳ")
            self.assertEqual(ws["J6"].value, "Xuất trong kỳ")
            self.assertEqual(ws["K6"].value, "Tồn cuối")
            self.assertEqual(ws["D7"].value, "Viên")
            self.assertEqual(ws["F7"].value, "31-12-2028")
            self.assertEqual(ws["H7"].value, 100)
            self.assertEqual(ws["K7"].value, 85)

            total_row = 9
            self.assertEqual(ws[f"A{total_row}"].value, "TỔNG CỘNG")
            self.assertEqual(ws[f"H{total_row}"].value, 110)
            self.assertEqual(ws[f"I{total_row}"].value, 30)
            self.assertEqual(ws[f"J{total_row}"].value, 43)
            self.assertEqual(ws[f"K{total_row}"].value, 97)

            self.assertEqual(ws.freeze_panes, "A7")
            self.assertEqual(ws.page_setup.orientation, "landscape")
            self.assertEqual(ws.page_setup.fitToWidth, 1)
            self.assertEqual(ws.print_title_rows, "$5:$6")
            self.assertIn("Người lập", ws["A12"].value)
            self.assertIn("Thủ kho", ws["E12"].value)
            self.assertIn("Phụ trách đơn vị", ws["I12"].value)
            self.assertTrue(ws["C7"].alignment.wrap_text)

    def test_exporter_does_not_claim_statutory_form_number(self):
        app = _App()
        with tempfile.TemporaryDirectory() as td:
            path = os.path.join(td, "xnt.xlsx")
            app._write_xnt_operational_workbook(
                path,
                app.db.xnt_report("2026-07-01", "2026-07-31", "Tất cả"),
                "2026-07-01",
                "2026-07-31",
                "Tất cả",
            )
            ws = load_workbook(path).active
            visible_text = " ".join(
                str(cell.value or "")
                for row in ws.iter_rows()
                for cell in row
            )
            self.assertNotIn("S12-H", visible_text)
            self.assertNotIn("S21-H", visible_text)
            self.assertNotIn("S22-H", visible_text)
            self.assertNotIn("S23-H", visible_text)


if __name__ == "__main__":
    unittest.main()
