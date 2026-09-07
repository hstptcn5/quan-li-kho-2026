# -*- coding: utf-8 -*-
"""Printable Excel form for the operational XNT report.

This module owns presentation only. Quantities come unchanged from DB.xnt_report().
The workbook is intentionally labelled as an operational XNT report rather than
claiming a statutory Vietnamese accounting form number.
"""

from __future__ import annotations

import os
from tkinter import filedialog, messagebox

from date_utils import format_date_display


class XntExcelExportMixin:
    """Replace the raw-ish XNT workbook with a printable operational form."""

    def export_report_excel(self):
        start_s, end_s = self._date_range_from_entries(self.de_from, self.de_to) if hasattr(self, "de_from") and hasattr(self, "de_to") else ("", "")
        if not start_s or not end_s:
            messagebox.showwarning("Thiếu ngày", "Chọn đủ Từ ngày và Đến ngày")
            return

        fund_source = self.cmb_report_fund.get().strip() if hasattr(self, "cmb_report_fund") else "Tất cả"
        rows = self.db.xnt_report(start_s, end_s, fund_source)
        if not rows:
            messagebox.showinfo("Không có dữ liệu", "Không có dữ liệu Xuất - Nhập - Tồn trong khoảng đã chọn.")
            return

        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx")],
            initialfile=f"Bao_cao_XNT_{start_s}_{end_s}.xlsx",
        )
        if not path:
            return

        try:
            self._write_xnt_operational_workbook(path, rows, start_s, end_s, fund_source)
            try:
                os.startfile(path)
            except Exception:
                pass
            self.toast("Đã xuất báo cáo XNT Excel theo mẫu")
        except Exception as exc:
            messagebox.showerror("Lỗi", f"Không thể xuất Excel: {exc}")

    def _write_xnt_operational_workbook(self, path, rows, start_s, end_s, fund_source):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.page import PageMargins

        wb = Workbook()
        ws = wb.active
        ws.title = "Báo cáo XNT"
        ws.sheet_view.showGridLines = False

        # Resolve base units once without changing the report query/calculation.
        unit_rows = self.db.q("SELECT id, COALESCE(defaultUnit, '') AS defaultUnit FROM products")
        unit_by_product = {int(r["id"]): (r["defaultUnit"] or "") for r in unit_rows}

        last_col = 11
        last_letter = get_column_letter(last_col)

        # ---- Document header -------------------------------------------------
        ws.merge_cells(f"A1:{last_letter}1")
        ws["A1"] = "BÁO CÁO XUẤT - NHẬP - TỒN"
        ws["A1"].font = Font(name="Arial", size=16, bold=True)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 25

        ws.merge_cells(f"A2:{last_letter}2")
        ws["A2"] = f"Từ ngày {format_date_display(start_s)} đến ngày {format_date_display(end_s)}"
        ws["A2"].font = Font(name="Arial", size=10, italic=True)
        ws["A2"].alignment = Alignment(horizontal="center")

        ws.merge_cells(f"A3:{last_letter}3")
        ws["A3"] = f"Nguồn kinh phí: {fund_source or 'Tất cả'}"
        ws["A3"].font = Font(name="Arial", size=10)
        ws["A3"].alignment = Alignment(horizontal="center")

        # ---- Two-level table header -----------------------------------------
        header_top = 5
        header_bottom = 6
        static_headers = [
            ("A", "STT"), ("B", "Mã SP"), ("C", "Tên thuốc / vaccine / VTYT"),
            ("D", "ĐVT"), ("E", "Số lô"), ("F", "Hạn sử dụng"), ("G", "Nguồn kinh phí"),
        ]
        for col, label in static_headers:
            ws.merge_cells(f"{col}{header_top}:{col}{header_bottom}")
            ws[f"{col}{header_top}"] = label

        ws.merge_cells(f"H{header_top}:K{header_top}")
        ws[f"H{header_top}"] = "SỐ LƯỢNG"
        for col, label in zip("HIJK", ("Tồn đầu", "Nhập trong kỳ", "Xuất trong kỳ", "Tồn cuối")):
            ws[f"{col}{header_bottom}"] = label

        thin = Side(style="thin", color="808080")
        medium = Side(style="medium", color="404040")
        table_border = Border(left=thin, right=thin, top=thin, bottom=thin)
        header_fill = PatternFill("solid", fgColor="D9EAF7")
        total_fill = PatternFill("solid", fgColor="E2F0D9")

        for row in ws.iter_rows(min_row=header_top, max_row=header_bottom, min_col=1, max_col=last_col):
            for cell in row:
                cell.font = Font(name="Arial", size=9, bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = table_border
        ws.row_dimensions[header_top].height = 22
        ws.row_dimensions[header_bottom].height = 30

        # ---- Data ------------------------------------------------------------
        data_start = 7
        totals = [0.0, 0.0, 0.0, 0.0]
        for idx, r in enumerate(rows, start=1):
            row_no = data_start + idx - 1
            product_id = int(r["productId"])
            values = [
                idx,
                product_id,
                r["productName"] or "",
                unit_by_product.get(product_id, ""),
                r["lotNo"] or "",
                format_date_display(r["expiryDate"]) if r["expiryDate"] else "",
                r["fundSource"] or "",
                float(r["opening"] or 0),
                float(r["inbound"] or 0),
                float(r["outbound"] or 0),
                float(r["closing"] or 0),
            ]
            for col_idx, value in enumerate(values, start=1):
                cell = ws.cell(row=row_no, column=col_idx, value=value)
                cell.font = Font(name="Arial", size=9)
                cell.border = table_border
                cell.alignment = Alignment(
                    horizontal="left" if col_idx in (3, 7) else "center",
                    vertical="center",
                    wrap_text=col_idx in (3, 7),
                )
            for offset, value in enumerate(values[7:11]):
                totals[offset] += float(value)
            for col_idx in range(8, 12):
                ws.cell(row=row_no, column=col_idx).number_format = '#,##0.####'
            ws.row_dimensions[row_no].height = 30

        total_row = data_start + len(rows)
        ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=7)
        ws.cell(total_row, 1, "TỔNG CỘNG")
        ws.cell(total_row, 1).alignment = Alignment(horizontal="center", vertical="center")
        for i, total in enumerate(totals, start=8):
            ws.cell(total_row, i, total)
            ws.cell(total_row, i).number_format = '#,##0.####'
        for cell in ws[total_row]:
            cell.font = Font(name="Arial", size=9, bold=True)
            cell.fill = total_fill
            cell.border = Border(left=thin, right=thin, top=medium, bottom=medium)
            if cell.column >= 8:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[total_row].height = 24

        # ---- Signature block -------------------------------------------------
        signature_top = total_row + 3
        ws.merge_cells(start_row=signature_top, start_column=1, end_row=signature_top, end_column=3)
        ws.merge_cells(start_row=signature_top, start_column=5, end_row=signature_top, end_column=7)
        ws.merge_cells(start_row=signature_top, start_column=9, end_row=signature_top, end_column=11)
        for coord, text in ((f"A{signature_top}", "Người lập"), (f"E{signature_top}", "Thủ kho"), (f"I{signature_top}", "Phụ trách đơn vị")):
            ws[coord] = text
            ws[coord].font = Font(name="Arial", size=10, bold=True)
            ws[coord].alignment = Alignment(horizontal="center")

        signature_note_row = signature_top + 1
        for start_col in (1, 5, 9):
            ws.merge_cells(start_row=signature_note_row, start_column=start_col, end_row=signature_note_row, end_column=start_col + 2)
            cell = ws.cell(signature_note_row, start_col, "(Ký, ghi rõ họ tên)")
            cell.font = Font(name="Arial", size=8, italic=True)
            cell.alignment = Alignment(horizontal="center")

        # ---- Usability + printing -------------------------------------------
        widths = {"A": 6, "B": 10, "C": 34, "D": 10, "E": 14, "F": 14, "G": 20, "H": 12, "I": 13, "J": 13, "K": 12}
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        ws.freeze_panes = "A7"
        ws.auto_filter.ref = f"A{header_bottom}:{last_letter}{total_row}"
        ws.print_title_rows = f"{header_top}:{header_bottom}"
        ws.print_area = f"A1:{last_letter}{signature_note_row + 4}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.5, bottom=0.5, header=0.2, footer=0.25)
        ws.oddFooter.center.text = "Trang &[Page]/&[Pages]"
        ws.oddFooter.center.size = 8
        ws.oddFooter.right.text = "Báo cáo XNT"
        ws.oddFooter.right.size = 8

        wb.save(path)
